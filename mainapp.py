import customtkinter as ctk
from tkinter import filedialog, messagebox
import fitz  # PyMuPDF
from docx import Document
import json
from googletrans import Translator
from gtts import gTTS
import pygame

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Smart Translator")
        self.root.geometry("900x700")
        self.root.resizable(True, True)

        ctk.set_appearance_mode("System")  # Modes: "System" (default), "Dark", "Light"
        ctk.set_default_color_theme("dark-blue")  # Themes: "blue" (default), "green", "dark-blue"

        self.selected_color = "System"  # Default color mode
        self.uploaded_text = ""  # Variable to store uploaded text

        self.create_sign_in_frame()

    def create_sign_in_frame(self):
        self.sign_in_frame = ctk.CTkFrame(self.root)
        self.sign_in_frame.pack(fill='both', expand=True, padx=30, pady=30)

        ctk.CTkLabel(self.sign_in_frame, text="Sign In", font=ctk.CTkFont(size=24, weight="bold")).pack(pady=20)

        ctk.CTkLabel(self.sign_in_frame, text="Username:", font=ctk.CTkFont(size=16)).pack(pady=5)
        self.username_entry = ctk.CTkEntry(self.sign_in_frame, width=300, height=40, font=ctk.CTkFont(size=14))
        self.username_entry.pack(pady=5)

        ctk.CTkLabel(self.sign_in_frame, text="Password:", font=ctk.CTkFont(size=16)).pack(pady=5)
        self.password_entry = ctk.CTkEntry(self.sign_in_frame, show="*", width=300, height=40, font=ctk.CTkFont(size=14))
        self.password_entry.pack(pady=5)

        button_frame = ctk.CTkFrame(self.sign_in_frame)
        button_frame.pack(pady=30)

        ctk.CTkButton(button_frame, text="MODE", command=self.toggle_mode, width=140, height=40, font=ctk.CTkFont(size=14)).pack(side='left', padx=10)
        ctk.CTkButton(button_frame, text="Sign In", command=self.sign_in, width=140, height=40, font=ctk.CTkFont(size=14)).pack(side='right', padx=10)
        
        self.status_label = ctk.CTkLabel(self.sign_in_frame, text="", text_color="red", font=ctk.CTkFont(size=14))
        self.status_label.pack(pady=10)

    def toggle_mode(self):
        # Toggle between dark mode and light mode
        new_mode = "Dark" if ctk.get_appearance_mode() == "Light" else "Light"
        ctk.set_appearance_mode(new_mode)

    def sign_in(self):
        username = self.username_entry.get()
        password = self.password_entry.get()

        if not username or not password:
            self.status_label.configure(text="Username and password are required.")
        elif username == "admin" and password == "password":  # Replace with actual logic
            self.status_label.configure(text="")
            self.sign_in_frame.pack_forget()
            self.create_file_upload_frame()  # Directly go to the file upload page
        else:
            self.status_label.configure(text="Invalid username or password.")

    def create_file_upload_frame(self):
        self.file_upload_frame = ctk.CTkFrame(self.root)
        self.file_upload_frame.pack(fill='both', expand=True, padx=30, pady=30)

        ctk.CTkLabel(self.file_upload_frame, text="File Upload", font=ctk.CTkFont(size=24, weight="bold")).pack(pady=20)

        self.back_button = ctk.CTkButton(self.file_upload_frame, text="Back", command=self.go_back, width=120, height=40, font=ctk.CTkFont(size=14))
        self.back_button.pack(pady=10)

        ctk.CTkButton(self.file_upload_frame, text="Upload File", command=self.upload_file, width=160, height=40, font=ctk.CTkFont(size=16)).pack(pady=20)

        self.text_area = ctk.CTkTextbox(self.file_upload_frame, wrap='word', height=200, font=ctk.CTkFont(size=14))
        self.text_area.pack(padx=20, pady=20, fill='both', expand=True)

        ctk.CTkButton(self.file_upload_frame, text="Next", command=self.create_translation_frame, width=140, height=40, font=ctk.CTkFont(size=16)).pack(pady=20)

    def go_back(self):
        self.file_upload_frame.pack_forget()
        self.create_sign_in_frame()

    def upload_file(self):
        file_path = filedialog.askopenfilename(
            filetypes=[("All Files", "*.*"), ("PDF Files", "*.pdf"), ("DOCX Files", "*.docx"), ("JSON Files", "*.json")]
        )
        if not file_path:
            return

        file_type = file_path.split('.')[-1].lower()
        if file_type == 'pdf':
            self.read_pdf(file_path)
        elif file_type == 'docx':
            self.read_docx(file_path)
        elif file_type == 'json':
            self.read_json(file_path)
        else:
            messagebox.showerror("Unsupported File", "Unsupported file type")

    def read_pdf(self, file_path):
        text = ""
        try:
            pdf_document = fitz.open(file_path)
            for page_num in range(len(pdf_document)):
                page = pdf_document.load_page(page_num)
                text += page.get_text()
            pdf_document.close()
            self.text_area.delete(1.0, ctk.END)
            self.text_area.insert(ctk.END, text)
            self.uploaded_text = text  # Save the uploaded text for the translation frame
        except Exception as e:
            messagebox.showerror("Error", f"Error reading PDF file: {e}")

    def read_docx(self, file_path):
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
            self.text_area.delete(1.0, ctk.END)
            self.text_area.insert(ctk.END, text)
            self.uploaded_text = text  # Save the uploaded text for the translation frame
        except Exception as e:
            messagebox.showerror("Error", f"Error reading DOCX file: {e}")

    def read_json(self, file_path):
        data = {}
        try:
            with open(file_path, 'r') as f:
                data = json.load(f)
            self.text_area.delete(1.0, ctk.END)
            self.text_area.insert(ctk.END, json.dumps(data, indent=4))
            self.uploaded_text = json.dumps(data, indent=4)  # Save the uploaded text for the translation frame
        except Exception as e:
            messagebox.showerror("Error", f"Error reading JSON file: {e}")

    def create_translation_frame(self):
        self.file_upload_frame.pack_forget()

        self.translation_frame = ctk.CTkFrame(self.root)
        self.translation_frame.pack(fill='both', expand=True, padx=30, pady=30)

        self.text_frame = ctk.CTkFrame(self.translation_frame)
        self.text_frame.grid(row=0, column=0, padx=10, pady=10, sticky="nsew")

        # Display uploaded text in the translation frame as the initial input
        text_label = ctk.CTkLabel(self.text_frame, text="Uploaded Text:", font=ctk.CTkFont(size=16, weight="bold"))
        text_label.pack(anchor="w")

        self.text_entry = ctk.CTkTextbox(self.text_frame, wrap='word', font=ctk.CTkFont(size=14))
        self.text_entry.insert(ctk.END, self.uploaded_text)  # Set the uploaded text as the initial input
        self.text_entry.pack(fill='both', expand=True, pady=(0, 10))

        controls_frame = ctk.CTkFrame(self.translation_frame)
        controls_frame.grid(row=0, column=1, padx=10, pady=10, sticky="nsew")

        language_label = ctk.CTkLabel(controls_frame, text="Select Target Language:", font=ctk.CTkFont(size=16))
        language_label.pack(pady=5)

        self.language_combobox = ctk.CTkComboBox(controls_frame, values=['te', 'ta', 'hi'], font=ctk.CTkFont(size=14))
        self.language_combobox.set('te')  # Default language
        self.language_combobox.pack(pady=5)

        translate_button = ctk.CTkButton(controls_frame, text="Translate", command=self.on_translate, width=160, height=40, font=ctk.CTkFont(size=16))
        translate_button.pack(pady=10)

        translated_label = ctk.CTkLabel(controls_frame, text="Translated Text:", font=ctk.CTkFont(size=16))
        translated_label.pack(pady=5)

        self.translated_text_area = ctk.CTkTextbox(controls_frame, wrap='word', font=ctk.CTkFont(size=14))
        self.translated_text_area.pack(pady=5, fill='both', expand=True)

        # Add Pause and Resume buttons side by side
        audio_controls_frame = ctk.CTkFrame(controls_frame)
        audio_controls_frame.pack(pady=10)

        self.pause_button = ctk.CTkButton(audio_controls_frame, text="Pause", command=self.pause_audio, width=160, height=40, font=ctk.CTkFont(size=16))
        self.pause_button.pack(side='left', padx=10)

        self.resume_button = ctk.CTkButton(audio_controls_frame, text="Resume", command=self.resume_audio, width=160, height=40, font=ctk.CTkFont(size=16))
        self.resume_button.pack(side='right', padx=10)

        back_button = ctk.CTkButton(self.translation_frame, text="Back to Upload", command=self.go_back_to_upload, width=160, height=40, font=ctk.CTkFont(size=16))
        back_button.grid(row=1, column=0, columnspan=2, pady=20)

        self.translation_frame.grid_columnconfigure(0, weight=1)
        self.translation_frame.grid_columnconfigure(1, weight=1)
        self.translation_frame.grid_rowconfigure(0, weight=1)

    def on_translate(self):
        target_language = self.language_combobox.get()
        translator = Translator()
        translated_text = translator.translate(self.text_entry.get("1.0", ctk.END), dest=target_language).text
        self.translated_text_area.delete(1.0, ctk.END)
        self.translated_text_area.insert(ctk.END, translated_text)
        self.play_translation(translated_text)

    def play_translation(self, translated_text=None):
        if translated_text is None:
            translated_text = self.translated_text_area.get("1.0", ctk.END).strip()
        tts = gTTS(text=translated_text, lang=self.language_combobox.get())
        tts.save("translation.mp3")
        pygame.mixer.init()
        pygame.mixer.music.load("translation.mp3")
        pygame.mixer.music.play()

    def pause_audio(self):
        pygame.mixer.music.pause()

    def resume_audio(self):
        pygame.mixer.music.unpause()

    def go_back_to_upload(self):
        self.translation_frame.pack_forget()
        self.create_file_upload_frame()

# Main Execution
if __name__ == "__main__":
    root = ctk.CTk()
    app = App(root)
    root.mainloop()
